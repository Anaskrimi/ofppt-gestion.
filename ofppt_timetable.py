"""Logique de l'emploi du temps (Time Table), SANS interface graphique :
constantes, conflits, import des seances eNote, analyse de la masse horaire, export Excel, PDF.

Ce fichier doit rester dans le MEME dossier que app.py."""
import os
import re
import json
import gzip
import math
import unicodedata
from datetime import datetime, timedelta
import pandas as pd

from ofppt_core import get_db_connection, EMPLOI_FILE, clean_str, logger

# ------------------------------------------------------------------
# MODULE EMPLOI DU TEMPS (Time Table) - integre dans le logiciel
# Les donnees sont sauvegardees automatiquement dans emploi_data.json
# (et on peut aussi les enregistrer / ouvrir via le menu Fichier).
# ------------------------------------------------------------------
TT_DAYS = ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi"]
TT_Q, TT_START = 42, 480                      # 42 quarts d'heure a partir de 08:00 (-> 18:30)
TT_LW, TT_HH, TT_QW, TT_RH = 90, 34, 24, 66         # geometrie de la grille
TT_GREEN, TT_BLUE, TT_GREY, TT_ORANGE = "#068a4e", "#1d5a8c", "#8e99a4", "#e06f3f"
TT_ETAB = [("direction", "Direction"), ("complexe", "Complexe"), ("nom", "Etablissement"), ("abrev", "Abréviation"),
        ("code", "Code"), ("annee", "Année de formation"), ("debut", "Date début (AAAA-MM-JJ)"), ("fin", "Date fin (AAAA-MM-JJ)")]
TT_SCHEMA = {
    "formateurs": ("Les Formateurs", TT_GREEN, [("nom", "Nom", None), ("prenom", "Prénom", None), ("matricule", "Matricule", None),
                                             ("type", "Type", ["Permanent", "Vacataire"])]),
    "filieres": ("Les Filières", TT_BLUE, [("nom", "Nom", None), ("abrev", "Abréviation", None),
                ("niveau", "Niveau", ["Technicien Spécialisé", "Technicien", "Qualification", "Spécialisation"]),
                ("secteur", "Secteur", ["Bâtiment et Travaux Publics", "Gestion et Commerce", "Digital et IA", "Agriculture", "Industrie"]),
                ("type", "Type de formation", ["Diplômante", "Qualifiante"])]),
    "groupes": ("Les Groupes", TT_BLUE, [("nom", "Nom", None), ("filiere", "Filière", "@f"), ("annee", "Année", ["1", "2", "3"])]),
    "espaces": ("Les Espaces", TT_GREY, [("nom", "Nom", None), ("type", "Type", ["Salle", "Atelier"])]),
}
TT_KIND = {"esp": ("espaces", "espace", "Espace", TT_GREY), "grp": ("groupes", "groupe", "Groupe", TT_BLUE),
        "frm": ("formateurs", "formateur", "Formateur", TT_GREEN)}

# Couleur des seances = UNE COULEUR PAR SALLE / ATELIER (grille ET PDF).
# Chaque espace garde la meme couleur partout. Modifiez la palette si besoin.
TT_PALETTE = ["#1d5a8c", "#0e7c66", "#d98219", "#7c3aed", "#c0392b",
              "#0891b2", "#65a30d", "#db2777", "#8a5a2b", "#4f46e5", "#0d9488", "#b91c1c"]

def tt_color(D, espace_id):
    """Couleur d'une seance selon son espace (salle/atelier).
    La couleur personnalisee (D['couleurs']) prime sur la palette par defaut."""
    c = (D.get("couleurs") or {}).get(str(espace_id))
    if c:
        return c
    try:
        return TT_PALETTE[int(espace_id) % len(TT_PALETTE)]
    except (TypeError, ValueError):
        return TT_PALETTE[0]

def tt_rgb(html):
    return tuple(int(html[n:n + 2], 16) for n in (1, 3, 5))


def tt_read_file(path):
    """Lit un fichier Time Table (.tk ou .json) : JSON brut ou compresse (gzip), encodage UTF-8 / ANSI."""
    with open(path, "rb") as fh:
        raw = fh.read()
    if raw[:2] == b"\x1f\x8b":
        raw = gzip.decompress(raw)
    last = None
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            x = json.loads(raw.decode(enc))
            if isinstance(x, dict):
                return x
            last = ValueError("structure inattendue")
        except Exception as e:
            last = e
    raise last


def tt_norm(s):
    return " ".join(str(s or "").casefold().split())


def tt_absents(date_obj):
    """Noms (normalises) des agents absents a une date : maladie, formation (suivi d'absence + formations)."""
    out = set()
    try:
        with get_db_connection() as conn:
            cur = conn.cursor()
            cur.execute("""SELECT p.nom, p.prenom, a.date_debut, a.date_fin FROM absences a
                           JOIN personnel p ON a.personnel_id = p.id WHERE a.motif IN ('Maladie', 'Formation')
                           UNION ALL
                           SELECT p.nom, p.prenom, f.date_debut, f.date_fin FROM formations f
                           JOIN personnel p ON f.personnel_id = p.id""")
            rows = cur.fetchall()
    except Exception:
        return out
    d = date_obj.date() if isinstance(date_obj, datetime) else date_obj
    for nom, prenom, d1, d2 in rows:
        try:
            a = datetime.strptime(d1, "%d/%m/%Y").date()
            b = datetime.strptime(d2, "%d/%m/%Y").date()
        except Exception:
            continue
        if a <= d <= b:
            out.add(tt_norm(f"{nom} {prenom}"))
            out.add(tt_norm(f"{prenom} {nom}"))
    return out


def tt_qt(q):
    m = TT_START + q * 15
    return f"{m // 60:02d}:{m % 60:02d}"


def tt_dur_txt(q):
    m = q * 15
    return f"{m // 60} h" + (f" {m % 60:02d}" if m % 60 else "")


def tt_blank():
    return {"file": "Sans titre", "etab": {k: "" for k, _ in TT_ETAB}, "filieres": [], "groupes": [], "formateurs": [],
            "espaces": [], "seances": [], "n": 1, "v": 3}


def tt_migrate(d):                          # anciens fichiers : debut a 08:30, puis heures entieres
    if d.get("v") != 3:
        if d.get("v") != 2:
            for s in d["seances"]:
                s["start"] *= 4; s["dur"] *= 4
        for s in d["seances"]:
            s["start"] += 2
        d["v"] = 3
    return d


def tt_get(D, k, i):
    return next((x for x in D[k] if x["id"] == i), None)


def tt_fname(f):
    return f"{f['nom']} {f.get('prenom', '')}".strip()


def tt_conflict(D, s):
    if s["start"] < 0 or s["start"] + s["dur"] > TT_Q:
        return "La séance dépasse la plage 08:00 - 18:30."
    for t in D["seances"]:
        if t["id"] == s["id"] or t["day"] != s["day"] or t["start"] >= s["start"] + s["dur"] or s["start"] >= t["start"] + t["dur"]:
            continue
        if t["espace"] == s["espace"]: return "Cet espace est déjà occupé sur ce créneau."
        if t["groupe"] == s["groupe"]: return "Ce groupe a déjà une séance sur ce créneau."
        if s["formateur"] and t["formateur"] == s["formateur"]: return "Ce formateur a déjà une séance sur ce créneau."
    return None


def tt_seance_text(D, s, kind, maxlen=28):
    """Texte affiche dans une case de seance (selon qu'on regarde un espace, un groupe ou un formateur)."""
    g = tt_get(D, "groupes", s["groupe"]) or {"nom": "?", "filiere": 0}; fi = tt_get(D, "filieres", g["filiere"]) or {}
    gl = f"[{fi.get('abrev') or fi.get('nom', '')}] : {g['nom']}"
    fm = tt_get(D, "formateurs", s["formateur"]); fl = tt_fname(fm).upper() if fm else ""
    el = (tt_get(D, "espaces", s["espace"]) or {"nom": "?"})["nom"]
    txt = {"esp": f"{fl}\n{gl}", "grp": f"{fl}\n{el}", "frm": f"{gl}\n{el}"}[kind]
    mt = s.get("matiere") or ""
    if mt: txt += "\n" + (mt if len(mt) <= maxlen else mt[:maxlen - 1] + "…")
    return txt


def tt_pdf_emploi(D, items):
    """Genere un PDF (A4 paysage) : une page par emploi du temps.
    items = liste de (kind, id) avec kind dans {"grp", "frm", "esp"}. Retourne l'objet FPDF."""
    from fpdf import FPDF                                  # import tardif : module utilisable sans fpdf installe
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(False)
    X0, LW, Y0, HH, RH = 10, 24, 30, 9, 24
    QW = (297 - 2 * X0 - LW) / TT_Q
    GX1 = X0 + LW + TT_Q * QW
    etab = clean_str((D.get("etab") or {}).get("nom", ""))

    def fit(txt, w):
        t = clean_str(txt.replace("…", "..."))
        if pdf.get_string_width(t) > w:
            while t and pdf.get_string_width(t + "..") > w: t = t[:-1]
            t += ".."
        return t

    for kind, i in items:
        coll, fld, lab, col = TT_KIND[kind]; o = tt_get(D, coll, i)
        if not o: continue
        lst = [s for s in D["seances"] if s[fld] == i]
        rgb = tuple(int(col[n:n + 2], 16) for n in (1, 3, 5))
        pdf.add_page()
        if etab:
            pdf.set_font("Arial", "", 9); pdf.set_text_color(90, 90, 90); pdf.set_xy(X0, 8); pdf.cell(0, 5, etab)
        nom = tt_fname(o) if kind == "frm" else o["nom"]
        m = sum(s["dur"] for s in lst) * 15
        pdf.set_fill_color(*rgb); pdf.set_text_color(255, 255, 255); pdf.set_font("Arial", "B", 14)
        pdf.set_xy(X0, 14); pdf.cell(297 - 2 * X0 - 48, 11, clean_str(f"  {lab} : {nom}"), 0, 0, "L", True)
        pdf.cell(48, 11, f"{m // 60:02d} H {m % 60:02d} Min", 0, 0, "C", True)

        # en-tete des heures
        pdf.set_draw_color(0, 0, 0); pdf.set_line_width(0.4); pdf.set_text_color(0, 0, 0); pdf.set_font("Arial", "", 6.5)
        heads = [(0, 2)] + [(2 + 4 * n, 4) for n in range(10)]
        for q, ln in heads:
            x = X0 + LW + q * QW
            pdf.rect(x, Y0, ln * QW, HH)
            pdf.set_xy(x, Y0 + HH / 2 - 1.5); pdf.cell(ln * QW, 3, f"{tt_qt(q)}-{tt_qt(q + ln)}", 0, 0, "C")
        # jours + quadrillage
        pdf.set_font("Arial", "B", 9)
        for d, name in enumerate(TT_DAYS):
            y = Y0 + HH + d * RH
            pdf.set_line_width(0.4); pdf.rect(X0, y, LW, RH)
            pdf.set_xy(X0, y + RH / 2 - 2); pdf.cell(LW, 4, name, 0, 0, "C")
            pdf.set_line_width(0.2); pdf.set_draw_color(120, 120, 120); pdf.line(X0 + LW, y, GX1, y)
            pdf.set_draw_color(0, 0, 0)
        pdf.set_line_width(0.4); pdf.line(X0 + LW, Y0 + HH + 6 * RH, GX1, Y0 + HH + 6 * RH)
        for q in range(TT_Q + 1):
            hour = q == 0 or q % 4 == 2
            pdf.set_line_width(0.4 if hour else 0.1); pdf.set_draw_color(0 if hour else 170, 0 if hour else 170, 0 if hour else 170)
            pdf.line(X0 + LW + q * QW, Y0 + HH, X0 + LW + q * QW, Y0 + HH + 6 * RH)
        pdf.set_draw_color(0, 0, 0)

        # seances
        for s in lst:
            x1, y1 = X0 + LW + s["start"] * QW + 0.4, Y0 + HH + s["day"] * RH + 0.6
            w, h = s["dur"] * QW - 0.8, RH - 1.2
            pdf.set_fill_color(*tt_rgb(tt_color(D, s.get("espace"))))
            pdf.rect(x1, y1, w, h, "F")
            pdf.set_text_color(255, 255, 255); pdf.set_font("Arial", "B", 6.5)
            lines = tt_seance_text(D, s, kind).split("\n")[:max(1, int(h // 3.2))]
            yy = y1 + (h - len(lines) * 3.2) / 2
            for ln_txt in lines:
                pdf.set_xy(x1 + 0.5, yy); pdf.cell(w - 1, 3.2, fit(ln_txt, w - 1.5), 0, 0, "C"); yy += 3.2
        pdf.set_text_color(120, 120, 120); pdf.set_font("Arial", "I", 7); pdf.set_xy(X0, 200)
        pdf.cell(0, 4, clean_str("Édité le " + datetime.now().strftime("%d/%m/%Y %H:%M")))
    return pdf


# ------------------------------------------------------------------
# IMPORT DES SEANCES eNOTE (export Excel de la page "Seances realisees")
# ------------------------------------------------------------------
TT_ENOTE_ALIAS = {
    "date": ("date", "dateseance", "datedeseance"),
    "groupe": ("groupe", "groupes"),
    "module": ("module",),
    "elt": ("eltcompt", "eltcomp", "element"),
    "hd": ("hd", "hdebut", "heuredebut", "debut"),
    "hf": ("hf", "hfin", "heurefin", "fin"),
    "local": ("typedelocal", "typelocal", "local"),
    "statut": ("statut",),
    "formateur": ("formateur", "nomformateur"),
    "semaine": ("semaine",),
}


def tt_enote_key(s):
    s = unicodedata.normalize("NFD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]", "", s.casefold())


def tt_enote_val(v):
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except (TypeError, ValueError):
        pass
    return str(v).strip()


def tt_enote_hm(v):
    """'13:30' / heure Excel / datetime -> minutes depuis minuit (None si illisible)."""
    if v is None or (isinstance(v, float) and v != v):
        return None
    if hasattr(v, "hour") and hasattr(v, "minute"):
        return v.hour * 60 + v.minute
    if isinstance(v, (int, float)):
        return int(round(v * 1440)) if 0 <= v < 1 else None
    m = re.search(r"(\d{1,2})\s*[:hH]\s*(\d{2})", str(v))
    return int(m.group(1)) * 60 + int(m.group(2)) if m else None


def tt_enote_date(v):
    if v is None or (isinstance(v, float) and v != v):
        return None
    try:
        if hasattr(v, "date") and callable(v.date):
            return v.date()
        s = str(v).strip()
        m = re.match(r"(\d{4})-(\d{2})-(\d{2})", s)
        if m:
            return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3))).date()
        m = re.match(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})", s)
        if m:
            return datetime(int(m.group(3)), int(m.group(2)), int(m.group(1))).date()
    except (ValueError, OverflowError):
        pass
    return None


def tt_enote_load(path):
    """Lit un export Excel/CSV d'eNote et renvoie la liste des seances (dictionnaires)."""
    ext = os.path.splitext(path)[1].lower()
    df = None
    if ext == ".csv":
        for enc in ("utf-8-sig", "cp1252"):
            try:
                df = pd.read_csv(path, sep=None, engine="python", header=None, dtype=object, encoding=enc)
                break
            except Exception:
                df = None
    else:
        try:
            df = pd.read_excel(path, header=None, dtype=object)
        except Exception:
            try:                                    # certains exports .xls sont en fait du HTML
                df = pd.read_html(path, header=None)[0].astype(object)
            except Exception:
                df = None
    if df is None or df.empty:
        raise ValueError("fichier illisible (Excel ou CSV attendu ; pour un .xls : pip install xlrd, ou enregistrez-le en .xlsx)")
    hdr = None
    for i in range(min(len(df), 20)):
        keys = [tt_enote_key(x) for x in df.iloc[i].tolist()]
        if "date" in keys and "groupe" in keys:
            hdr = i
            break
    if hdr is None:
        raise ValueError("colonnes « Date » et « Groupe » introuvables")
    keys = [tt_enote_key(x) for x in df.iloc[hdr].tolist()]
    col = {}
    for name, aliases in TT_ENOTE_ALIAS.items():
        for j, k in enumerate(keys):
            if k in aliases:
                col[name] = j
                break
    for need in ("date", "groupe", "hd", "hf"):
        if need not in col:
            raise ValueError(f"colonne « {need} » introuvable")
    rows = []
    for _, r in df.iloc[hdr + 1:].iterrows():
        vals = r.tolist()
        get = lambda n: vals[col[n]] if n in col else None
        d = tt_enote_date(get("date"))
        groupe = tt_enote_val(get("groupe"))
        if d is None or not groupe:
            continue
        rows.append({"date": d, "groupe": groupe, "module": tt_enote_val(get("module")),
                     "elt": tt_enote_val(get("elt")), "hd": tt_enote_hm(get("hd")), "hf": tt_enote_hm(get("hf")),
                     "local": tt_enote_val(get("local")), "statut": tt_enote_val(get("statut")),
                     "formateur": tt_enote_val(get("formateur")), "semaine": tt_enote_val(get("semaine"))})
    return rows


def tt_enote_apply(D, uid, lots, lundi):
    """Place dans l'emploi du temps D les seances de la semaine commencant le `lundi`.
    lots = [(nom du formateur par defaut, [seances])]. Cree les groupes, formateurs et
    salles manquants. Renvoie un compte rendu."""
    rep = {"ajoutees": 0, "deja": 0, "conflits": [], "ignorees": [], "crees": [], "par_formateur": {}}

    def get_grp(nom):
        key = tt_norm(nom)
        g = next((x for x in D["groupes"] if tt_norm(x["nom"]) == key), None)
        if g:
            return g
        m = re.match(r"[A-Za-z]+", nom)
        pref = m.group(0).upper() if m else nom
        fi = next((x for x in D["filieres"] if tt_norm(x.get("abrev") or x["nom"]) == tt_norm(pref)), None)
        if not fi:
            fi = {"id": uid(), "nom": pref, "abrev": pref, "niveau": "Technicien Spécialisé", "secteur": "", "type": "Diplômante"}
            D["filieres"].append(fi)
            rep["crees"].append(f"Filière « {pref} » (à renommer dans Paramètres > Filières)")
        an = re.search(r"[123]", nom)
        g = {"id": uid(), "nom": nom, "filiere": fi["id"], "annee": an.group(0) if an else "1"}
        D["groupes"].append(g)
        rep["crees"].append(f"Groupe « {nom} »")
        return g

    def get_frm(nom):
        if not nom:
            return None
        toks = sorted(tt_norm(nom).split())
        for x in D["formateurs"]:
            if sorted(tt_norm(tt_fname(x)).split()) == toks:
                return x["id"]
        nm, pr = tt_split_name(nom)
        f = {"id": uid(), "nom": nm, "prenom": pr, "matricule": "", "type": "Permanent"}
        D["formateurs"].append(f)
        rep["crees"].append(f"Formateur « {nom} »")
        return f["id"]

    def occupe(eid, day, start, dur):
        return any(t["espace"] == eid and t["day"] == day and not (t["start"] >= start + dur or start >= t["start"] + t["dur"])
                   for t in D["seances"])

    pref_salle = {}

    def get_espace(gid, typ, day, start, dur):
        cands = ([pref_salle[gid]] if gid in pref_salle else []) + [e["id"] for e in D["espaces"] if e["type"] == typ]
        for eid in cands:
            if not occupe(eid, day, start, dur):
                return eid
        n = sum(1 for e in D["espaces"] if e["type"] == typ) + 1
        while any(tt_norm(e["nom"]) == tt_norm(f"{typ} {n}") for e in D["espaces"]):
            n += 1
        e = {"id": uid(), "nom": f"{typ} {n}", "type": typ}
        D["espaces"].append(e)
        rep["crees"].append(f"Espace « {e['nom']} » (salle non précisée dans eNote)")
        return e["id"]

    items = []
    for fdef, rows in lots:
        for r in rows:
            d = r["date"]
            if d - timedelta(days=d.weekday()) != lundi:
                continue
            lib = f"{d.strftime('%d/%m')} {r['groupe']}"
            if d.weekday() > 5:
                rep["ignorees"].append(f"{lib} : dimanche")
                continue
            if r["hd"] is None or r["hf"] is None or r["hf"] <= r["hd"]:
                rep["ignorees"].append(f"{lib} : horaires illisibles")
                continue
            start, dur = round((r["hd"] - TT_START) / 15), round((r["hf"] - r["hd"]) / 15)
            if start < 0 or start + dur > TT_Q:
                rep["ignorees"].append(f"{lib} : en dehors de 08:00 - 18:30")
                continue
            items.append((d.weekday(), start, dur, fdef or r.get("formateur"), r))
    items.sort(key=lambda t: (t[0], t[1]))

    for day, start, dur, fnom, r in items:
        g = get_grp(r["groupe"])
        fid = get_frm(fnom)
        lib = f"{TT_DAYS[day]} {tt_qt(start)}-{tt_qt(start + dur)} {g['nom']}"
        if any(t["groupe"] == g["id"] and t["day"] == day and t["start"] == start and t["dur"] == dur
               and t.get("formateur") == fid for t in D["seances"]):
            rep["deja"] += 1
            continue
        s = {"id": -1, "groupe": g["id"], "espace": None, "formateur": fid, "day": day, "start": start, "dur": dur}
        m = tt_conflict(D, s)
        if m:
            rep["conflits"].append(f"{lib} ({fnom}) : {m}")
            continue
        typ = "Atelier" if "atelier" in tt_norm(r["local"]) else "Salle"
        s["espace"] = get_espace(g["id"], typ, day, start, dur)
        s["id"] = uid()
        s["matiere"] = r["module"] or r["elt"]
        D["seances"].append(s)
        pref_salle[g["id"]] = s["espace"]
        rep["ajoutees"] += 1
        k = fnom or "(sans formateur)"
        rep["par_formateur"][k] = rep["par_formateur"].get(k, 0) + 1
    return rep


def tt_split_name(full):
    """'AKOUAOUACH MOHAMED' -> ('AKOUAOUACH', 'MOHAMED') : le dernier mot est le prenom."""
    p = str(full or "").split()
    return (" ".join(p[:-1]), p[-1]) if len(p) > 1 else (" ".join(p), "")


def tt_load_emploi():
    """Charge les donnees du module Emploi du temps (sauvegarde automatique) ; None si absentes."""
    try:
        with open(EMPLOI_FILE, encoding="utf-8") as fh:
            x = json.load(fh)
        d = {**tt_blank(), **x}
        d["v"] = x.get("v")
        return tt_migrate(d)
    except Exception:
        return None


def tt_analyse(D):
    """Masse horaire de chaque formateur : quarts d'heure par jour (lundi..samedi) et total de la semaine."""
    out = []
    for f in D["formateurs"]:
        ss = sorted((s for s in D["seances"] if s.get("formateur") == f["id"]), key=lambda s: (s["day"], s["start"]))
        jours = [sum(s["dur"] for s in ss if s["day"] == d) for d in range(6)]
        out.append({"nom": tt_fname(f), "type": f.get("type", ""), "jours": jours, "total": sum(jours), "seances": ss})
    return out


def tt_analyse_excel(path, D, res):
    """Exporte l'analyse : une feuille Resume + une feuille (tableau) par formateur."""
    from openpyxl.styles import Font, PatternFill
    h = lambda q: round(q / 4, 2)
    cols = ["Formateur", "Type"] + [f"{j} (h)" for j in TT_DAYS] + ["Total semaine (h)", "Séances"]
    rows = [[a["nom"], a["type"]] + [h(q) for q in a["jours"]] + [h(a["total"]), len(a["seances"])] for a in res if a["total"]]
    tot = [sum(a["jours"][d] for a in res) for d in range(6)]
    rows.append(["TOTAL", ""] + [h(q) for q in tot] + [h(sum(tot)), sum(len(a["seances"]) for a in res)])
    used = {"résumé"}
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        pd.DataFrame(rows, columns=cols).to_excel(w, sheet_name="Résumé", index=False)
        for a in res:
            if not a["total"]:
                continue
            name = re.sub(r"[\[\]:*?/\\]", " ", a["nom"]).strip()[:28] or "Formateur"
            base, n = name, 2
            while name.casefold() in used:
                name = f"{base[:25]} {n}"; n += 1
            used.add(name.casefold())
            det = []
            for s in a["seances"]:
                g = tt_get(D, "groupes", s["groupe"]) or {"nom": "?"}
                e = tt_get(D, "espaces", s["espace"]) or {"nom": "?"}
                det.append([TT_DAYS[s["day"]], tt_qt(s["start"]), tt_qt(s["start"] + s["dur"]), g["nom"],
                            s.get("matiere", ""), e["nom"], h(s["dur"])])
            det.append(["TOTAL SEMAINE", "", "", "", "", "", h(a["total"])])
            pd.DataFrame(det, columns=["Jour", "Début", "Fin", "Groupe", "Module", "Espace", "Durée (h)"]).to_excel(
                w, sheet_name=name, index=False)
            par_jour = [[TT_DAYS[d], h(a["jours"][d])] for d in range(6)] + [["Total semaine", h(a["total"])]]
            pd.DataFrame(par_jour, columns=["Jour", "Heures"]).to_excel(w, sheet_name=name, index=False, startcol=8)
        for ws in w.book.worksheets:
            for c in ws[1]:
                if c.value is not None:
                    c.font = Font(bold=True, color="FFFFFF")
                    c.fill = PatternFill("solid", fgColor="1F538D")
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(40, max(len(str(c.value or "")) for c in col) + 2)
